from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select, extract, func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.log_cuti import LogCuti
from app.models.log_cuti_date import LogCutiDate
from app.models.user import User
from app.services.cuti_service import hitung_cuti_terpakai
from app.services.date_format_service import BULAN_INDONESIA, format_tanggal_grouped


def _format_dates(logs: list[LogCuti]) -> str:
    if not logs:
        return "-"

    all_dates = []
    for log in logs:
        for log_date in log.tanggal_list:
            all_dates.append(log_date.tanggal)

    return format_tanggal_grouped(all_dates)


## log cuti final milik user pada tahun tertentu
## (tahun = min(tanggal_list) per log, pengganti extract kolom tanggal_mulai lama)
async def get_approved_logs_for_year(user_id: int, year: int, db: AsyncSession) -> list[LogCuti]:
    approved_statuses = ["disetujui_hr", "disetujui_direktur", "cuti_bersama"]

    min_tgl = (
        select(LogCutiDate.id_log_cuti, func.min(LogCutiDate.tanggal).label("mulai"))
        .group_by(LogCutiDate.id_log_cuti)
        .subquery()
    )

    result_logs = await db.execute(
        select(LogCuti).options(
            selectinload(LogCuti.tanggal_list)
        )
        .join(min_tgl, min_tgl.c.id_log_cuti == LogCuti.id_log_cuti)
        .where(
            LogCuti.id_user == user_id,
            LogCuti.status.in_(approved_statuses),
            extract("year", min_tgl.c.mulai) == year,
        )
    )
    return result_logs.scalars().unique().all()


async def export_cuti_excel(year: int, db: AsyncSession) -> bytes:
    approved_statuses = ["disetujui_hr", "disetujui_direktur", "cuti_bersama"]

    result_users = await db.execute(
        select(User)
        .where(User.role.in_(["karyawan", "pm", "hr_manager", "staff_hr"]))
        .order_by(User.nama.asc())
    )
    users = result_users.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Cuti"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    ws.merge_cells("A1:G1")
    ws["A1"] = f"Rekap Data Cuti Karyawan - Periode {year}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["No", "Nama Karyawan", "Total Cuti (hari)", "Total Cuti Diambil", "Sisa Cuti", "Keterangan (Tanggal Cuti)"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    row = 4
    no = 1

    for user in users:
        logs = await get_approved_logs_for_year(user.id_user, year, db)
        cuti_terpakai = await hitung_cuti_terpakai(user.id_user, db)
        keterangan = _format_dates(logs)

        ws.cell(row=row, column=1, value=no).border = border
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")

        ws.cell(row=row, column=2, value=user.nama).border = border
        ws.cell(row=row, column=3, value=user.total_cuti).border = border
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=4, value=cuti_terpakai).border = border
        ws.cell(row=row, column=4).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=5, value=user.sisa_cuti).border = border
        ws.cell(row=row, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=6, value=keterangan).border = border

        no += 1
        row += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 50

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return buffer.getvalue()
